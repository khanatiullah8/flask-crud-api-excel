from flask import Flask, jsonify, request, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_marshmallow import Marshmallow
from datetime import datetime
import openpyxl

# app init
app = Flask(__name__)

# config
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///students.db"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

# instance
db = SQLAlchemy(app)
ma = Marshmallow(app)

# model
class Student(db.Model):
  id = db.Column(db.Integer, primary_key=True)
  name = db.Column(db.String(100), nullable=False)
  age = db.Column(db.Integer, nullable=False)
  gender = db.Column(db.String(10), nullable=False)
  email = db.Column(db.String(100), nullable=False, unique=True)
  phone = db.Column(db.String(15), nullable=False)
  course = db.Column(db.String(100), nullable=False)
  enrollment_date = db.Column(db.Date, nullable=False)
  city = db.Column(db.String(100), nullable=False)
  grade = db.Column(db.String(10), nullable=False)
  
# schema
class StudentSchema(ma.Schema):
  class Meta:
    model = Student
    fields = ["id","name","age","gender","email","phone","course","enrollment_date","city","grade"]
    load_instance = True
    
# schema instance
student_schema = StudentSchema()
students_schema = StudentSchema(many=True)


# utility function
def str_to_date(value):
  return datetime.strptime(value, "%Y-%m-%d").date()


# routes
@app.route("/")
def home():
  return jsonify({"message":"Home page"})

# get all students
@app.route("/students", methods=["GET"])
def get_all_students():
  students = Student.query.all()
  return students_schema.dump(students)

# get student
@app.route("/students/<int:student_id>", methods=["GET"])
def get_student(student_id):
  student = Student.query.get(student_id)
  return student_schema.dump(student)

# add student
@app.route("/students", methods=["POST"])
def add_students():
  temp_data = request.get_json()
  data = student_schema.load(temp_data)
  data["enrollment_date"] = str_to_date(data["enrollment_date"])
  student = Student(**data)
  db.session.add(student)
  db.session.commit()
  return jsonify({"message":"student added successfully"})

# update student
@app.route("/students/<int:student_id>", methods=["PUT"])
def update_student(student_id):
  data = request.get_json()
  student = Student.query.get(student_id)
  student.name = data.get("name") or student.name
  student.age = data.get("age") or student.age
  student.gender = data.get("gender") or student.gender
  student.email = data.get("email") or student.email
  student.phone = data.get("phone") or student.phone
  student.course = data.get("course") or student.course
  if data.get("enrollment_date"):
    student.enrollment_date = str_to_date(data["enrollment_date"])
  student.city = data.get("city") or student.city
  student.grade = data.get("grade") or student.grade
  db.session.commit()
  return jsonify({"message":"student updated successfully"})

# delete student
@app.route("/students/<int:student_id>", methods=["DELETE"])
def delete_student(student_id):
  student = Student.query.get(student_id)
  db.session.delete(student)
  db.session.commit()
  return jsonify({"message":"student deleted successfully"})
  
# add/import students from excel file
@app.route("/students/import", methods=["POST"])
def import_students():
  file = request.files["file"]
  wb = openpyxl.load_workbook(file)
  sheet = wb.active
  for row in sheet.iter_rows(min_row=2, values_only=True):
    if all(row):
      if type(row[6]).__name__ == "datetime":
        dt = row[6].strftime("%Y-%m-%d")
        e_date = str_to_date(dt)
      elif type(row[6]).__name__ == "str":
        e_date = str_to_date(row[6])
      student = Student(name=row[0],age=row[1],gender=row[2],email=row[3],phone=str(int(row[4])),course=row[5],enrollment_date=e_date,city=row[7],grade=row[8])
      db.session.add(student)
    db.session.commit()
  return jsonify({"message":"excel file imported successfully"})

# export students as excel file
@app.route("/students/export", methods=["GET"])
def export_students():
  students = Student.query.all()
  wb = openpyxl.Workbook()
  sheet = wb.active
  sheet.append(["Name","Age","Gender","Email","Phone","Course","Enrollment_date","City","Grade"])
  for s in students:
    sheet.append([s.name, s.age, s.gender, s.email, s.phone, s.course, (s.enrollment_date).strftime("%d/%m/%Y"), s.city, s.grade])
  filename = "Students.xlsx"
  wb.save(filename=filename)
  return send_file(filename, as_attachment=True)

# app run
if __name__ == "__main__":
  with app.app_context():
    db.create_all()
    
  app.run(debug=True)